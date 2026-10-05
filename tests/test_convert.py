"""Tests for the optional LibreOffice conversion (mocked soffice — not installed here)."""

from pathlib import Path

import pytest

from course_assistant.ingest.convert import convert_decks, find_soffice


def test_find_soffice_returns_none_when_missing(monkeypatch):
    monkeypatch.setattr("course_assistant.ingest.convert.shutil.which", lambda name: None)
    assert find_soffice() is None


def test_convert_decks_raises_friendly_error_without_soffice(tmp_path, monkeypatch):
    monkeypatch.setattr("course_assistant.ingest.convert.shutil.which", lambda name: None)
    (tmp_path / "deck.pptx").write_bytes(b"x")
    with pytest.raises(RuntimeError, match="brew install --cask libreoffice"):
        convert_decks(tmp_path)


def test_convert_decks_checks_page_count_matches_slides(tmp_path, monkeypatch):
    # Fake soffice present; fake subprocess + pymupdf page count.
    monkeypatch.setattr("course_assistant.ingest.convert.shutil.which", lambda name: "/fake/soffice")
    monkeypatch.setattr(
        "course_assistant.ingest.convert.subprocess.run",
        lambda *a, **k: None,
    )
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for n in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"S{n}"
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    calls = {"n": 0}

    def fake_page_count(pdf):
        calls["n"] += 1
        return 3  # matches the 3 slides

    monkeypatch.setattr("course_assistant.ingest.convert._page_count", fake_page_count)
    results = convert_decks(tmp_path, tmp_path / "out")
    assert results == [("deck.pptx", 3, 3, True)]
