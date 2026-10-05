"""Tests for the parse layer — all formats generated on the fly."""

import base64
import io
from pathlib import Path

import pytest

from course_assistant.ingest.formats import UnsupportedFormatError, parse_document

PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def make_pptx(path: Path, slides: int = 2, with_picture: bool = True) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for n in range(1, slides + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {n} title"
        body = slide.placeholders[1]
        body.text = f"Body text of slide {n}, mentioning embedded concepts."
        if with_picture:
            slide.shapes.add_picture(io.BytesIO(PNG_1PX), Inches(0.5), Inches(0.5))
    prs.save(str(path))
    return path


def make_pdf(path: Path, pages: int = 2) -> Path:
    import fitz

    doc = fitz.open()
    for n in range(1, pages + 1):
        page = doc.new_page()
        page.insert_text((72, 72), f"Hello from PDF page {n}")
    doc.save(str(path))
    doc.close()
    return path


def make_docx(path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading(f"Docx Section", level=1)
    doc.add_paragraph("A regular paragraph with meaningful text.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    doc.add_picture(io.BytesIO(PNG_1PX), width=Inches(1))
    doc.save(str(path))
    return path


@pytest.fixture
def images_root(tmp_path: Path) -> Path:
    root = tmp_path / "images"
    root.mkdir()
    return root


def test_sha256_identity_and_dedupe_key(tmp_path, images_root):
    f = make_pptx(tmp_path / "deck.pptx")
    d1 = parse_document(f, images_root)
    d2 = parse_document(f, images_root)
    assert d1.doc_id == d2.doc_id == d1.doc_id  # same file -> same doc_id
    f2 = make_pptx(tmp_path / "other.pptx", slides=3)
    assert parse_document(f2, images_root).doc_id != d1.doc_id


def test_pptx_slides_text_title_and_images(tmp_path, images_root):
    d = parse_document(make_pptx(tmp_path / "deck.pptx"), images_root)
    assert d.format == ".pptx"
    assert d.total_pages == 2
    assert d.pages[0].section == "Slide 1 title"
    assert "Body text of slide 1" in d.pages[0].text
    assert len(d.pages[0].image_paths) == 1
    assert Path(d.pages[0].image_paths[0]).exists()
    assert d.pages[0].image_paths[0].startswith(str(images_root / d.doc_id))


def test_pptx_without_pictures_has_no_images(tmp_path, images_root):
    d = parse_document(make_pptx(tmp_path / "deck.pptx", with_picture=False), images_root)
    assert d.pages[0].image_paths == []


def test_pdf_pages_text_and_rendered_images(tmp_path, images_root):
    d = parse_document(make_pdf(tmp_path / "doc.pdf"), images_root)
    assert d.format == ".pdf"
    assert d.total_pages == 2
    assert "Hello from PDF page 2" in d.pages[1].text
    assert len(d.pages[0].image_paths) == 1  # rendered page PNG
    assert Path(d.pages[0].image_paths[0]).exists()


def test_docx_text_table_and_inline_image(tmp_path, images_root):
    d = parse_document(make_docx(tmp_path / "doc.docx"), images_root)
    text = d.pages[0].text
    assert "regular paragraph" in text
    assert "Cell A | Cell B" in text
    assert d.pages[0].section == "Docx Section"
    assert len(d.pages[0].image_paths) == 1
    assert Path(d.pages[0].image_paths[0]).exists()


def test_markdown_and_txt_single_page(tmp_path, images_root):
    md = tmp_path / "notes.md"
    md.write_text("# Heading\nContent line.")
    d = parse_document(md, images_root)
    assert d.total_pages == 1
    assert "Content line" in d.pages[0].text

    txt = tmp_path / "plain.txt"
    txt.write_text("just some text")
    assert "just some text" in parse_document(txt, images_root).pages[0].text


def test_html_strips_tags(tmp_path, images_root):
    html = tmp_path / "page.html"
    html.write_text("<html><body><h1>Title</h1><p>Para <b>one</b></p></body></html>")
    d = parse_document(html, images_root)
    assert "Title" in d.pages[0].text
    assert "Para one" in d.pages[0].text
    assert "<b>" not in d.pages[0].text


def test_unsupported_format_message_lists_accepted(tmp_path, images_root):
    bad = tmp_path / "file.xyz"
    bad.write_text("nope")
    with pytest.raises(UnsupportedFormatError) as ei:
        parse_document(bad, images_root)
    msg = str(ei.value)
    assert ".pdf" in msg and ".pptx" in msg and ".docx" in msg


def test_parse_is_deterministic(tmp_path, images_root):
    f = make_pptx(tmp_path / "deck.pptx")
    d1 = parse_document(f, images_root)
    d2 = parse_document(f, images_root)
    assert [p.text for p in d1.pages] == [p.text for p in d2.pages]
    assert [p.section for p in d1.pages] == [p.section for p in d2.pages]
